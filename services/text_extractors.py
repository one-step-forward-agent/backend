import csv
import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader


# Upper bounds for untrusted documents: page count, extracted text size, table rows
MAX_PDF_PAGES = 200
MAX_TEXT_CHARS = 200_000
MAX_ROWS = 2000
# A PDF with less text than this is a scan: its pictures are read instead
SCAN_TEXT_CHARS = 40
MAX_SCAN_IMAGES = 5

DOCUMENT_SUFFIXES = (".pdf", ".docx", ".xlsx", ".csv", ".tsv", ".txt", ".md", ".ics")


def extract_pdf(path: str) -> str:
    reader = PdfReader(path)
    parts, size = [], 0
    for page in reader.pages[:MAX_PDF_PAGES]:
        text = page.extract_text() or ""
        parts.append(text)
        size += len(text)
        if size >= MAX_TEXT_CHARS:
            break
    return "\n".join(parts).strip()[:MAX_TEXT_CHARS]


def pdf_images(path: str) -> list[tuple[bytes, str]]:
    """The pictures of a scanned PDF's first pages, as (data, mime) for the vision model."""
    found = []
    for page in PdfReader(path).pages[:MAX_SCAN_IMAGES]:
        for image in page.images:
            suffix = Path(image.name).suffix.lower()
            mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(suffix)
            if mime:
                found.append((image.data, mime))
            if len(found) >= MAX_SCAN_IMAGES:
                return found
    return found


def extract_docx(path: str) -> str:
    """Paragraphs and tables: a timetable in Word is usually a table."""
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows[:MAX_ROWS]:
            cells = list(dict.fromkeys(cell.text.strip() for cell in row.cells))
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()[:MAX_TEXT_CHARS]


def extract_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts, size = [], 0
    try:
        for sheet in workbook.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(max_row=MAX_ROWS, values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    line = " | ".join(cells).rstrip(" |")
                    parts.append(line)
                    size += len(line)
                if size >= MAX_TEXT_CHARS:
                    break
    finally:
        workbook.close()
    return "\n".join(parts).strip()[:MAX_TEXT_CHARS]


def decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_table(path: str, delimiter: str) -> str:
    rows = csv.reader(io.StringIO(decode(Path(path).read_bytes()[: MAX_TEXT_CHARS * 2])), delimiter=delimiter)
    lines = [" | ".join(cell.strip() for cell in row) for _, row in zip(range(MAX_ROWS), rows) if any(cell.strip() for cell in row)]
    return "\n".join(lines).strip()[:MAX_TEXT_CHARS]


def extract_document(path: str) -> str:
    suffix = Path(path).suffix.lower()
    if suffix not in DOCUMENT_SUFFIXES:
        raise ValueError("Такой файл не прочитать: пришлите PDF, Word, Excel, CSV, текст или фото")
    try:
        if suffix == ".pdf":
            return extract_pdf(path)
        if suffix == ".docx":
            return extract_docx(path)
        if suffix == ".xlsx":
            return extract_xlsx(path)
        if suffix in (".csv", ".tsv"):
            return extract_table(path, "\t" if suffix == ".tsv" else ",")
        return decode(Path(path).read_bytes()[: MAX_TEXT_CHARS * 2]).strip()[:MAX_TEXT_CHARS]
    except Exception as error:  # malformed or hostile file: report it, don't crash
        raise ValueError("Не удалось прочитать документ") from error
